"""Import a Plan from a workbook, replacing the current plan. Includes the
legacy trade-batch-format migration described in feature.md §9. Port of
importFromWorkbook() (Renovation_Planner_v2.html:1658-1772)."""

from __future__ import annotations

import datetime
import re
from pathlib import Path

from openpyxl import load_workbook

from renovator.domain.models import (
    BuyStatus,
    LineType,
    Phase,
    Plan,
    ProjectSettings,
    Rate,
    Stage,
    TaskLine,
    TaskStatus,
    WeekendPolicy,
    new_id,
)
from renovator.domain.seed import (
    CATEGORY_DURATION,
    default_phases,
    default_rates,
    default_rooms,
    default_stages,
)
from renovator.engine.tasks import derive_task_ids
from renovator.tools.errors import ValidationError

_STAGE_NUM_PREFIX = re.compile(r"^\s*\d+\s*[·.\-]\s*")


def _sheet_rows(wb, name: str) -> list[dict]:
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h) if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(v is None for v in r):
            continue
        out.append({headers[i]: ("" if i >= len(r) or r[i] is None else r[i]) for i in range(len(headers))})
    return out


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _strip_stage_num(s: str) -> str:
    return _STAGE_NUM_PREFIX.sub("", s or "").strip()


def _num(v, default: float = 0) -> float:
    try:
        if v is None or v == "":
            return default
        return float(v)
    except (TypeError, ValueError):
        return default


def import_plan_from_excel(path: Path) -> Plan:
    wb = load_workbook(path, data_only=True)

    items_sheet_name = "Rooms & Items" if "Rooms & Items" in wb.sheetnames else (wb.sheetnames[0] if wb.sheetnames else None)
    if not items_sheet_name:
        raise ValidationError('No "Rooms & Items" sheet found — nothing to import.')
    item_rows = _sheet_rows(wb, items_sheet_name)
    if not item_rows:
        raise ValidationError("That sheet has no rows to import.")

    headers = list(item_rows[0].keys())
    has_stage_col = "Stage" in headers
    has_batch_col = "Trade Batch" in headers
    is_old = not has_stage_col

    # ---- rates -------------------------------------------------------
    new_rates: list[Rate] = []
    rate_key_by_label: dict[str, str] = {}
    for r in _sheet_rows(wb, "Rate Card"):
        label = _s(r.get("Rate"))
        if not label:
            continue
        key = new_id("rate")
        new_rates.append(Rate(key=key, label=label, unit=_s(r.get("Unit")) or "sqft", value=_num(r.get("Rate (INR)"))))
        rate_key_by_label[label] = key
    if not new_rates:
        new_rates = default_rates()

    # ---- stage label per row (Renovation_Planner_v2.html:1683-1685) --
    def row_stage(row: dict) -> str:
        if has_stage_col:
            return _s(row.get("Stage"))
        if has_batch_col and _s(row.get("Trade Batch")):
            return _s(row.get("Trade Batch"))
        return _strip_stage_num(_s(row.get("Phase")))

    order_source = _sheet_rows(wb, "Schedule") or item_rows

    new_stages: list[Stage] = []
    stage_id_by_label: dict[str, int] = {}
    for r in _sheet_rows(wb, "Stages"):
        label = _s(r.get("Stage"))
        if not label or label in stage_id_by_label:
            continue
        sid = len(new_stages) + 1
        new_stages.append(Stage(id=sid, label=label))
        stage_id_by_label[label] = sid
    for row in order_source:
        label = row_stage(row)
        if not label or label in stage_id_by_label:
            continue
        sid = len(new_stages) + 1
        new_stages.append(Stage(id=sid, label=label))
        stage_id_by_label[label] = sid
    if not new_stages:
        new_stages = default_stages()

    # ---- phases (Renovation_Planner_v2.html:1700-1711) ----------------
    new_phases: list[Phase] = []
    phase_id_by_label: dict[str, int] = {}
    phases_rows = _sheet_rows(wb, "Phases")
    if phases_rows:
        for r in phases_rows:
            label = _s(r.get("Phase"))
            if not label or label in phase_id_by_label:
                continue
            pid = len(new_phases) + 1
            new_phases.append(Phase(id=pid, label=label))
            phase_id_by_label[label] = pid
    elif not is_old:
        for row in order_source:
            label = _s(row.get("Phase"))
            if not label or label in phase_id_by_label:
                continue
            pid = len(new_phases) + 1
            new_phases.append(Phase(id=pid, label=label))
            phase_id_by_label[label] = pid
    if not new_phases:
        new_phases = default_phases()
    phase1 = new_phases[0].id
    phase_last = new_phases[-1].id

    # ---- items ----------------------------------------------------------
    new_rooms: list[str] = []
    new_items: list[TaskLine] = []
    lookup_by_tag: dict[str, str] = {}
    dep_text_by_item_id: dict[str, str] = {}

    for row in item_rows:
        room = _s(row.get("Room")) or "Unassigned"
        if room not in new_rooms:
            new_rooms.append(room)
        name = _s(row.get("Item")) or "Untitled item"
        line_type_raw = row.get("Line Type")
        line_type = line_type_raw if line_type_raw in (m.value for m in LineType) else LineType.INCLUSIVE.value
        is_rate = _s(row.get("Costing Method")).lower().startswith("rate")
        est_cost = _num(row.get("Est. Cost (INR)"))
        rate_key = rate_key_by_label.get(_s(row.get("Rate Used"))) if is_rate else None
        category = _s(row.get("Category")) or "Other"
        mandatory = _s(row.get("Mandatory/Optional") or "Mandatory").lower() != "optional"

        if is_old:
            item_phase = phase1 if mandatory else phase_last
        else:
            item_phase = phase_id_by_label.get(_s(row.get("Phase")), phase1)

        stage_label = row_stage(row)
        stage_id = stage_id_by_label.get(stage_label) if stage_label else None

        status_raw = row.get("Status")
        status = status_raw if status_raw in (s.value for s in TaskStatus) else TaskStatus.NOT_STARTED.value

        dur_raw = row.get("Duration (days)")
        if dur_raw in (None, "") :
            dur_raw = row.get("Duration Days")
        try:
            duration_days = int(float(dur_raw)) if dur_raw not in (None, "") else CATEGORY_DURATION.get(category, 1)
        except (TypeError, ValueError):
            duration_days = CATEGORY_DURATION.get(category, 1)

        buy_raw = _s(row.get("Buy Status"))
        buy_status = buy_raw if buy_raw in (b.value for b in BuyStatus) else BuyStatus.TO_BUY.value

        actual_raw = row.get("Actual Cost (INR)")
        actual_cost = None if actual_raw in (None, "") else _num(actual_raw)

        item = TaskLine(
            id=new_id("item"),
            task_id=_s(row.get("Task ID")) or None,
            room=room,
            name=name,
            category=category,
            line_type=line_type,
            worker_type=_s(row.get("Worker/Trade")),
            mandatory=mandatory,
            phase=item_phase,
            stage_id=stage_id,
            depends_on=[],
            cost_method="rate" if (is_rate and rate_key) else "custom",
            rate_key=rate_key if (is_rate and rate_key) else (new_rates[0].key if new_rates else None),
            qty=_num(row.get("Quantity")) if is_rate else 0,
            custom_amount=0 if (is_rate and rate_key) else est_cost,
            duration_days=duration_days,
            start_override=_s(row.get("Fixed Start")) or None,
            buy_status=buy_status,
            status=status,
            actual_cost=actual_cost,
            notes=_s(row.get("Notes")),
        )
        new_items.append(item)
        lookup_by_tag[f"{room} - {name} [{line_type}]"] = item.id
        dep_text_by_item_id[item.id] = _s(row.get("Depends On"))

    for item in new_items:
        text = dep_text_by_item_id.get(item.id, "")
        if not text:
            continue
        for tag in (t.strip() for t in text.split(";")):
            if tag and tag in lookup_by_tag:
                item.depends_on.append(lookup_by_tag[tag])

    derive_task_ids(new_items)

    # ---- project settings ----------------------------------------------
    project_start = None
    weekend_policy = WeekendPolicy.NONE
    sequence_stages = True
    blocked_dates: list[str] = []
    for r in _sheet_rows(wb, "Project"):
        key = _s(r.get("Setting")).lower()
        value = _s(r.get("Value"))
        if "start" in key and re.match(r"^\d{4}-\d{2}-\d{2}$", value):
            project_start = value
        if "weekend" in key:
            # "Saturdays only" (new) vs. the legacy Yes/No toggle a workbook
            # exported before weekend_policy existed would still have.
            if "saturday" in value.lower():
                weekend_policy = WeekendPolicy.SATURDAYS
            elif value[:1].lower() == "y":
                weekend_policy = WeekendPolicy.ALL
            else:
                weekend_policy = WeekendPolicy.NONE
        if "cascade" in key:
            sequence_stages = value[:1].lower() == "y"
        if "blocked" in key:
            blocked_dates = sorted(
                d for d in (part.strip() for part in value.split(",")) if re.match(r"^\d{4}-\d{2}-\d{2}$", d)
            )

    return Plan(
        settings=ProjectSettings(
            project_start=project_start or datetime.date.today().isoformat(),
            weekend_policy=weekend_policy,
            sequence_stages=sequence_stages,
            blocked_dates=blocked_dates,
        ),
        rooms=new_rooms or default_rooms(),
        rates=new_rates,
        phases=new_phases,
        stages=new_stages,
        items=new_items,
    )
